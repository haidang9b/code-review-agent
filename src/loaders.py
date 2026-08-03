"""Knowledge-document loading for ``.md``, ``.pdf`` and ``.docx`` files.

Loads with ``pypdf`` and ``python-docx`` directly (rather than
``langchain-community`` loaders) to avoid pulling in ``langchain-community``
as a dependency.

Every produced :class:`Document` carries the metadata the reviewer needs in
order to *cite* a rule rather than assert one:

``source``    file path relative to the knowledge root
``doc_type``  the folder the file was found in (``standards``, ``security``, …)
``tier``      priority tier derived from ``doc_type`` (see :mod:`src.config`)
``format``    ``markdown`` | ``pdf`` | ``docx`` — chooses the splitting strategy
``page``      1-indexed page number (PDF only)
``heading``   nearest section heading (DOCX only; Markdown headings are added
              later by the Markdown-aware splitter)
"""

from __future__ import annotations

import re
from pathlib import Path

from langchain_core.documents import Document
from pypdf import PdfReader

from .config import DOC_TYPE_TIERS, TIER_CODING_STANDARD, resolve_knowledge_path

SUPPORTED_SUFFIXES = {".md", ".markdown", ".pdf", ".docx"}

# Word writes lock files as ~$name.docx while a document is open.
_SKIP_PREFIXES = ("~$", ".")


class UnsupportedDocumentError(ValueError):
    """Raised when a file extension has no loader."""


def _clean(text: str) -> str:
    """Collapse repeated whitespace while preserving paragraph breaks."""
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def doc_type_for(path: Path, root: Path) -> str:
    """Derive ``doc_type`` from the first folder under the knowledge root.

    The folder name is the authority — never an LLM guess — so a document's
    priority tier is deterministic and auditable. Files sitting directly in the
    root, or in an unrecognised folder, become ``other``.
    """
    try:
        relative = path.resolve().relative_to(root.resolve())
    except ValueError:
        return "other"

    if len(relative.parts) < 2:
        return "other"

    folder = relative.parts[0].lower()
    return folder if folder in DOC_TYPE_TIERS else "other"


def tier_for(doc_type: str) -> int:
    """Priority tier for a ``doc_type``, defaulting to the standards tier."""
    return DOC_TYPE_TIERS.get(doc_type, TIER_CODING_STANDARD)


def _base_metadata(path: Path, root: Path) -> dict:
    doc_type = doc_type_for(path, root)
    try:
        source = path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        source = path.name
    return {"source": source, "doc_type": doc_type, "tier": tier_for(doc_type)}


# ---- Per-format loaders ---------------------------------------------------


def load_markdown(path: Path, root: Path) -> list[Document]:
    """Load a Markdown file as a single document.

    Kept whole on purpose: :func:`src.splitters.split_knowledge_documents`
    splits it by heading so each chunk knows which section it came from.
    """
    text = _clean(path.read_text(encoding="utf-8", errors="replace"))
    if not text:
        return []
    return [Document(page_content=text, metadata={**_base_metadata(path, root), "format": "markdown"})]


def load_pdf(path: Path, root: Path) -> list[Document]:
    """Load a PDF into one document per page, so findings can cite a page."""
    reader = PdfReader(str(path))
    base = {**_base_metadata(path, root), "format": "pdf"}

    documents: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if not text:
            continue
        documents.append(Document(page_content=text, metadata={**base, "page": index}))
    return documents


def _iter_docx_blocks(document):
    """Yield paragraphs and tables of a DOCX body in document order.

    ``python-docx`` exposes ``paragraphs`` and ``tables`` as separate flat
    lists, which loses their relative order. Standards documents routinely put
    the actual rule in a table directly under its heading, so order matters.
    """
    from docx.document import Document as _DocxBody
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body if isinstance(document, _DocxBody) else document._tc
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _render_table(table) -> str:
    """Render a DOCX table as pipe-delimited rows, readable as evidence."""
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def load_docx(path: Path, root: Path) -> list[Document]:
    """Load a Word document as one document per heading section.

    Paragraphs *and* table cells are walked in document order; heading styles
    start a new section so each chunk can be cited by its heading.
    """
    from docx import Document as read_docx
    from docx.table import Table

    document = read_docx(str(path))
    base = {**_base_metadata(path, root), "format": "docx"}

    sections: list[tuple[str, list[str]]] = [("", [])]
    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            rendered = _render_table(block)
            if rendered:
                sections[-1][1].append(rendered)
            continue

        text = block.text.strip()
        if not text:
            continue

        style = (block.style.name or "") if block.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            sections.append((text, []))
        else:
            sections[-1][1].append(text)

    documents: list[Document] = []
    for heading, lines in sections:
        body = _clean("\n\n".join(lines))
        if not body:
            continue
        content = f"{heading}\n\n{body}" if heading else body
        metadata = {**base, "heading": heading} if heading else dict(base)
        documents.append(Document(page_content=content, metadata=metadata))
    return documents


_LOADERS = {
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".pdf": load_pdf,
    ".docx": load_docx,
}


def load_document(path: str | Path, root: str | Path | None = None) -> list[Document]:
    """Load a single knowledge file, dispatching on its extension."""
    path = Path(path)
    root = Path(root) if root else resolve_knowledge_path()
    if not path.exists():
        raise FileNotFoundError(f"Knowledge document not found: {path}")

    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedDocumentError(
            f"No loader for '{path.suffix}'. Supported: {sorted(SUPPORTED_SUFFIXES)}"
        )
    return loader(path, root)


def discover_documents(root: str | Path | None = None) -> list[Path]:
    """List every supported knowledge file under ``root``, sorted for stability."""
    root = resolve_knowledge_path(root)
    if not root.is_dir():
        raise FileNotFoundError(
            f"Knowledge folder not found: {root}. Create it and add your coding "
            f"standards, ADRs and design documents, or set KNOWLEDGE_PATH."
        )

    return sorted(
        p
        for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        and not p.name.startswith(_SKIP_PREFIXES)
    )


def load_knowledge(root: str | Path | None = None) -> tuple[list[Document], list[str]]:
    """Load every supported document under the knowledge root.

    Returns the loaded documents plus a list of human-readable problems (a
    corrupt PDF, an unreadable DOCX). Failures are reported rather than raised
    so one bad file cannot block ingesting the rest of the corpus.
    """
    root = resolve_knowledge_path(root)
    documents: list[Document] = []
    problems: list[str] = []

    for path in discover_documents(root):
        try:
            loaded = load_document(path, root)
        except Exception as exc:  # noqa: BLE001 - surfaced to the operator, not swallowed
            problems.append(f"{path.name}: {type(exc).__name__}: {exc}")
            continue
        if not loaded:
            problems.append(f"{path.name}: no extractable text")
            continue
        documents.extend(loaded)

    return documents, problems
